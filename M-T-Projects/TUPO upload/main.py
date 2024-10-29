import  openpyxl, os, re, sys
from pdfminer.high_level import extract_pages, extract_text

from helpers import *

def main():
     # file to produce the invoice results 
    writePath = r'C:\Users\MOHIBIM\OneDrive - Ventia\Documents\M&T Finance\TUPO Output.xlsx'
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    # User provides the location of the folders where invoices are
    folLoc = r'C:\Users\MOHIBIM\OneDrive - Ventia\\Documents\M&T Finance\TU PO\\'

     # Walk through the folder and find the pdf files
    rowNum = 1
    for dirpath, direnames, files in os.walk(folLoc):
        for file in files: 
            print(file)
            filename = os.fsdecode(file)

             # Looks through all pdf files
            if filename.endswith(".pdf"):
                newFilePath = os.path.join(dirpath,filename)
                pdfText = extract_text(newFilePath)

                # Look for each type of AS based on the set rules of wiritin AS 6 digits, 7 digits, M6digits
                asNums = [(pdfText[x.start():x.end()]).strip() for x in re.finditer(r"2[0-9]{6}|1[0-9]{6}|4[0-9]{6}|M[0-9]{6}|1[0-9]{5}|2[0-9]{5}", pdfText)]
                PONums = [(pdfText[x.end():x.end()+5]).strip() for x in re.finditer(r"PO000", pdfText)]
                valueNums = [(pdfText[x.end():x.end()+11]).strip() for x in re.finditer(r"Line Amount", pdfText)]
                print(asNums, PONums,valueNums)
            else:
                print(f"{filename} is not a pdf")
            
            # Print to the excel sheet in the form of the TUPO tracker
            sheet.cell(row= rowNum,column=1).value = PONums[0]
            sheet.cell(row= rowNum,column=2).value = asNums[0]
            sheet.cell(row= rowNum,column=3).value =  valueNums[0]
            rowNum += 1
    
    workbook.save(filename=writePath)
main()